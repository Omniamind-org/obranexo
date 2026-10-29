"""Kit da proposta (PT-1.8, fase B1): montado da análise e do catálogo, sem IA, sem preço e sem envio."""

from __future__ import annotations

import io
from dataclasses import replace

import pytest
from fastapi.testclient import TestClient

from obranexo.catalog.atributos import unidade
from obranexo.catalog.loader import carregar_catalogo_json, carregar_catalogo_tabela
from obranexo.config import config
from obranexo.models import ResultadoItem
from obranexo.proposta import OFERTAVEIS, PREENCHER, Situacao, kit_markdown, kit_planilha, montar_kit, valor_tecnico

CFG = replace(config, llm_provider=None, offline=True)


@pytest.fixture(scope="module")
def empresa():
    return carregar_catalogo_json(config.catalogo_demo)


def _analise(caso, empresa):
    from obranexo.agents.pipeline import executar_caso

    return executar_caso(config.dir_casos / caso, empresa, cfg=CFG, salvar=False)


@pytest.fixture(scope="module")
def goianira(empresa):
    return _analise("goianira_pe038_2026", empresa)


@pytest.fixture(scope="module")
def aruana(empresa):
    return _analise("aruana_pp005_2026", empresa)


def test_goianira_itens_e_produtos(goianira, empresa):
    kit = montar_kit(goianira, empresa)
    sit = {i.numero: (i.situacao, i.produto.sku if i.produto else None) for i in kit.itens}
    assert sit == {
        1: (Situacao.OFERTAR, "LPL-100-50K"), 2: (Situacao.OFERTAR, "LPL-100-50K"),
        3: (Situacao.OFERTAR, "LPL-50-60K"), 4: (Situacao.OFERTAR, "LPL-50-60K"),
        5: (Situacao.OFERTAR_COM_RESSALVA, "RF-MULTI"), 6: (Situacao.OFERTAR_COM_RESSALVA, "RF-MULTI"),
    }
    item1 = kit.itens[0]
    assert (item1.produto.marca, item1.produto.modelo) == ("LUMIDEMO", "VIA 100 50K")
    assert kit.resumo["ofertaveis"] == 6 and kit.resumo["valor_estimado_ofertavel"] == pytest.approx(1_671_180.0)


def test_produto_incompativel_nunca_e_ofertado(aruana, empresa):
    kit = montar_kit(aruana, empresa)
    item1 = kit.itens[0]
    assert item1.situacao == Situacao.NAO_OFERTAR and item1.produto is None and item1.descricao_ofertada is None
    assert "não atende" in item1.motivo
    melhor = {}
    ordem = {ResultadoItem.COMPATIVEL: 3, ResultadoItem.PARCIAL: 2, ResultadoItem.INDETERMINADO: 1, ResultadoItem.INCOMPATIVEL: 0}
    for c in aruana.correspondencias:
        if c.item_id not in melhor or ordem[c.resultado] > ordem[melhor[c.item_id]]:
            melhor[c.item_id] = c.resultado
    for i in kit.itens:
        if i.situacao in OFERTAVEIS:
            assert melhor[i.item_id] in {ResultadoItem.COMPATIVEL, ResultadoItem.PARCIAL}


def test_descricao_so_com_fatos_do_catalogo(goianira, empresa):
    kit = montar_kit(goianira, empresa)
    item = kit.itens[0]
    produto = next(p for p in empresa.produtos if p.sku == item.produto.sku)
    for e in item.exigencias:
        no_texto = f"{e.rotulo}: {e.ofertado}" in item.descricao_ofertada
        if e.resultado.value == "ATENDE" and e.atributo in produto.atributos and e.atributo != "prazo_entrega_dias":
            assert no_texto, e.atributo
            assert e.ofertado == valor_tecnico(produto.atributos[e.atributo].valor, unidade(e.atributo))
        elif e.resultado.value != "ATENDE":
            assert not no_texto, e.atributo
    assert "Prazo de entrega" not in item.descricao_ofertada  # condição comercial fica nas condições


def test_condicoes_tem_prova_conferida(goianira, empresa):
    kit = montar_kit(goianira, empresa)
    rotulos = {c.rotulo: c for c in kit.condicoes}
    assert {"Fim do recebimento de propostas", "Prazo de entrega exigido", "Garantia mínima exigida", "Amostra"} <= set(rotulos)
    prazo = rotulos["Prazo de entrega exigido"]
    assert prazo.valor == "máximo 10 dias" and goianira.evidencia(prazo.evidencia_id).verificada


def test_produto_sem_marca_pede_preenchimento(goianira, empresa):
    sem_marca = empresa.model_copy(update={"produtos": [p.model_copy(update={"marca": None, "modelo": None}) for p in empresa.produtos]})
    kit = montar_kit(goianira, sem_marca)
    assert any("marca ou modelo" in a for a in kit.avisos)
    item = kit.itens[0]
    assert f"marca {PREENCHER}" in item.descricao_ofertada and any("marca e modelo" in p for p in item.pendencias)


def test_exportacoes(goianira, empresa):
    from openpyxl import load_workbook

    kit = montar_kit(goianira, empresa)
    md = kit_markdown(kit, goianira)
    assert "## Prazos e condições" in md and "### Item 1 — Ofertar" in md and "LUMIDEMO" in md
    assert "Sem preço" in md and "não envia propostas" in md
    wb = load_workbook(io.BytesIO(kit_planilha(kit)))
    ws = wb["Proposta"]
    cab = [c.value for c in ws[1]]
    assert cab[:3] == ["Item", "Situação", "Descrição do item (edital)"] and "Marca" in cab and "Modelo" in cab
    assert ws.max_row == 1 + len(kit.itens)
    assert ws.cell(row=2, column=cab.index("Marca") + 1).value == "LUMIDEMO"
    assert "Prazos e condições" in wb.sheetnames


def _docx(kit, empresa):
    from docx import Document

    from obranexo.proposta import proposta_docx

    doc = Document(io.BytesIO(proposta_docx(kit, empresa)))
    return doc, "\n".join(p.text for p in doc.paragraphs)


def test_proposta_de_precos_docx(goianira, empresa):
    doc, texto = _docx(montar_kit(goianira, empresa), empresa)
    assert "PROPOSTA DE PREÇOS" in texto and "Distribuidora Demo Elétrica Ltda. (fictícia)" in texto and "00.000.000/0001-00" in texto
    assert "Prazo de entrega: até 10 dias, conforme o edital" in texto
    assert "Garantia: conforme o edital (mínimo de 5 anos)" in texto
    assert f"Validade da proposta: {PREENCHER}" in texto  # o edital não foi lido para isso: não afirmamos
    assert "CONFERÊNCIA INTERNA — REMOVA ESTA PÁGINA ANTES DE ENVIAR" in texto
    tabela = doc.tables[0]
    assert len(tabela.rows) == 1 + 6 + 1  # cabeçalho, 6 itens ofertáveis, valor global
    item1 = [c.text for c in tabela.rows[1].cells]
    assert item1[0] == "1" and "Marca: LUMIDEMO" in item1[2] and "Modelo: VIA 100 50K" in item1[2]
    assert item1[1].startswith("Luminária pública LED 100 W 5000 K.") and "Fluxo luminoso: 14.500 lm" in item1[1]
    assert item1[5] == PREENCHER and item1[6] == PREENCHER  # preço é decisão do fornecedor


def test_proposta_docx_nao_oferta_incompativel(aruana, empresa):
    doc, texto = _docx(montar_kit(aruana, empresa), empresa)
    numeros = [r.cells[0].text for r in doc.tables[0].rows[1:-1]]
    assert numeros == ["2", "3"]  # item 1 não atende: fica fora da tabela
    assert "Item 1 (Não ofertar)" in texto  # e aparece na conferência interna, com o motivo


def test_proposta_docx_sem_dados_do_fornecedor(goianira, empresa):
    vazia = empresa.model_copy(update={"cnpj": None, "razao_social": None, "endereco": None, "representante": None, "dados_bancarios": None})
    _, texto = _docx(montar_kit(goianira, vazia), vazia)
    assert f"CNPJ: {PREENCHER}" in texto and f"Endereço: {PREENCHER}" in texto and f"Dados bancários: {PREENCHER}" in texto


def test_dados_do_fornecedor_persistem(tmp_path, monkeypatch, empresa):
    from obranexo.storage import store

    monkeypatch.setattr(config, "dir_dados", tmp_path)
    assert empresa.razao_social and empresa.cnpj == "00.000.000/0001-00"
    store.salvar_catalogo(empresa)
    relida = carregar_catalogo_json(store.caminho_catalogo_ativo())
    assert (relida.razao_social, relida.representante, relida.dados_bancarios) == (empresa.razao_social, empresa.representante, empresa.dados_bancarios)


def test_catalogo_em_planilha_le_marca_e_modelo(tmp_path):
    csv = tmp_path / "cat.csv"
    csv.write_text("sku;nome;categoria;fabricante;marca;modelo;potencia_w\nX1;Luminária;luminaria_publica_led;Fab Ltda;MARCA;MOD-1;100\n", encoding="utf-8")
    p = carregar_catalogo_tabela(csv).produtos[0]
    assert (p.fabricante, p.marca, p.modelo) == ("Fab Ltda", "MARCA", "MOD-1")


def test_importacao_de_ficha_traz_marca_e_modelo():
    from obranexo.catalog.importar import FichaExtraida, ProdutoFicha, importar_ficha

    class _LLM:
        disponivel = True

        def extrair(self, *_a, **_k):
            return FichaExtraida(produtos=[ProdutoFicha(nome="Luminária X 100 W", categoria="luminaria_publica_led",
                                                        fabricante=" Fab Ltda ", marca="MARCA", modelo="X-100")])

    produtos, _ = importar_ficha("ficha técnica " * 10, _LLM(), "modelo")
    assert (produtos[0].fabricante, produtos[0].marca, produtos[0].modelo) == ("Fab Ltda", "MARCA", "X-100")


@pytest.fixture()
def client(tmp_path, monkeypatch, entrar):
    monkeypatch.setattr(config, "dir_analises", tmp_path / "analises")
    monkeypatch.setattr(config, "dir_dados", tmp_path)
    monkeypatch.setattr(config, "llm_provider", None)
    from obranexo.api import main

    main._estado.clear()
    with TestClient(main.app) as c:
        c.usuario = entrar(c)
        yield c
    main._estado.clear()


def test_endpoints_do_kit(client, goianira):
    from obranexo import contas
    from obranexo.storage import store

    with contas.na_empresa(client.usuario.empresa_id):
        store.salvar(goianira)
        store.salvar(goianira.model_copy(update={"id": "com-erro", "status": "erro"}))
    base = f"/api/analises/{goianira.id}/proposta"
    r = client.get(base)
    assert r.status_code == 200 and r.json()["resumo"]["ofertaveis"] == 6
    assert client.get(base + ".md").text.startswith("# Kit da proposta")
    x = client.get(base + ".xlsx")
    assert x.status_code == 200 and "spreadsheetml" in x.headers["content-type"] and "attachment" in x.headers["content-disposition"]
    w = client.get(base + ".docx")
    assert w.status_code == 200 and "wordprocessingml" in w.headers["content-type"] and "proposta_de_precos" in w.headers["content-disposition"]
    assert client.get("/api/analises/nao-existe/proposta").status_code == 404
    assert client.get("/api/analises/com-erro/proposta").status_code == 409
