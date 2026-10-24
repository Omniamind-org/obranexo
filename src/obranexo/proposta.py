"""Kit da proposta: o que o fornecedor precisa para disputar a licitação, montado a partir da análise.

Sem IA nesta etapa (PT-1.8, fase B1). Tudo sai de fatos que o sistema já conferiu:
  - o produto ofertado e seus atributos vêm do catálogo do fornecedor ativo;
  - a conformidade vem do motor de regras (ATENDE / NÃO ATENDE / INDETERMINADO);
  - cada exigência e cada prazo apontam para o trecho do edital que os sustenta.

Regras fixas: produto INCOMPATÍVEL nunca é ofertado; item PARCIAL ou INDETERMINADO entra com o que falta
resolver; documento só aparece como "declarado" se o catálogo declarar; o kit não sugere preço (o valor
estimado pelo órgão é só referência) e não envia nada — quem revisa e envia é o fornecedor.
"""

from __future__ import annotations

import io
from datetime import datetime
from enum import StrEnum
from typing import Any

from pydantic import BaseModel, Field

from obranexo import __version__
from obranexo.catalog.atributos import natureza, rotulo, unidade
from obranexo.models import Analise, Correspondencia, DetalheRequisito, Empresa, ItemDemanda, Produto, ResultadoItem, ResultadoRequisito

PREENCHER = "[preencher]"
_ORDEM = {ResultadoItem.COMPATIVEL: 3, ResultadoItem.PARCIAL: 2, ResultadoItem.INDETERMINADO: 1, ResultadoItem.INCOMPATIVEL: 0}
_OPERADOR = {">=": "mínimo", "<=": "máximo", "entre": "entre"}
_COMERCIAIS = {"prazo_entrega_dias"}


class Situacao(StrEnum):
    OFERTAR = "OFERTAR"
    OFERTAR_COM_RESSALVA = "OFERTAR_COM_RESSALVA"
    RESOLVER_ANTES = "RESOLVER_ANTES"
    NAO_OFERTAR = "NAO_OFERTAR"
    FORA_DO_CATALOGO = "FORA_DO_CATALOGO"


ROTULO_SITUACAO = {
    Situacao.OFERTAR: "Ofertar",
    Situacao.OFERTAR_COM_RESSALVA: "Ofertar com ressalva",
    Situacao.RESOLVER_ANTES: "Resolver antes de ofertar",
    Situacao.NAO_OFERTAR: "Não ofertar",
    Situacao.FORA_DO_CATALOGO: "Fora do catálogo",
}
OFERTAVEIS = {Situacao.OFERTAR, Situacao.OFERTAR_COM_RESSALVA}


class ProdutoOfertado(BaseModel):
    sku: str
    nome: str
    fabricante: str | None = None
    marca: str | None = None
    modelo: str | None = None


class Exigencia(BaseModel):
    atributo: str
    rotulo: str
    obrigatorio: bool
    exigido: str
    ofertado: str
    resultado: ResultadoRequisito
    justificativa: str
    evidencia_id: str | None = None


class DocumentoTecnico(BaseModel):
    nome: str
    situacao: str  # declarado | falta | operacional | a_confirmar | recomendado
    detalhe: str
    exigido: bool
    evidencia_id: str | None = None


class ItemProposta(BaseModel):
    item_id: str
    numero: int
    descricao_orgao: str
    quantidade: float | None = None
    unidade: str | None = None
    valor_unit_estimado: float | None = None
    valor_total_estimado: float | None = None
    situacao: Situacao
    rotulo_situacao: str
    motivo: str
    produto: ProdutoOfertado | None = None
    descricao_ofertada: str | None = None
    exigencias: list[Exigencia] = Field(default_factory=list)
    pendencias: list[str] = Field(default_factory=list)
    documentos: list[DocumentoTecnico] = Field(default_factory=list)


class Condicao(BaseModel):
    rotulo: str
    valor: str
    fonte: str
    evidencia_id: str | None = None


class KitProposta(BaseModel):
    analise_id: str
    gerado_em: datetime
    versao: str = __version__
    contratacao: dict[str, Any]
    fornecedor: dict[str, Any]
    resumo: dict[str, Any]
    condicoes: list[Condicao]
    itens: list[ItemProposta]
    habilitacao: dict[str, Any]
    avisos: list[str]


# ---------------------------------------------------------------------------------------------- formatação


def _num(v: float) -> str:
    """Número em pt-BR: 14500 -> 14.500; 0.98 -> 0,98; até duas casas decimais."""
    inteiro, dec = f"{float(v):,.2f}".split(".")
    dec = dec.rstrip("0")
    inteiro = inteiro.replace(",", ".")
    return f"{inteiro},{dec}" if dec else inteiro


def valor_tecnico(v: Any, und: str | None = None) -> str:
    if v is None or v == "":
        return "—"
    sufixo = f" {und}" if und else ""
    if isinstance(v, bool):
        return "sim" if v else "não"
    if isinstance(v, list | tuple) and len(v) == 2:
        return f"{_num(v[0])} a {_num(v[1])}{sufixo}"
    if isinstance(v, int | float):
        return f"{_num(v)}{sufixo}"
    return str(v)


def _exigido(req_operador: str | None, valor: Any, atributo: str) -> str:
    if req_operador == "verdadeiro" or isinstance(valor, bool):
        return "exigido"
    texto = valor_tecnico(valor, unidade(atributo))
    prefixo = _OPERADOR.get(req_operador or "")
    return f"{prefixo} {texto}" if prefixo and prefixo != "entre" else texto


# ---------------------------------------------------------------------------------------------- montagem


def _melhor(correspondencias: list[Correspondencia]) -> Correspondencia | None:
    return max(correspondencias, key=lambda c: (_ORDEM[c.resultado], c.score), default=None)


def _tecnicos(exigencias: list[Exigencia]) -> list[str]:
    """Atributos do catálogo que o motor de regras confirmou, na ordem do edital. Prazo de entrega é condição
    comercial: fica nas condições, não na especificação do produto."""
    return [f"{e.rotulo}: {e.ofertado}" for e in exigencias
            if e.resultado == ResultadoRequisito.ATENDE and natureza(e.atributo) == "produto" and e.atributo not in _COMERCIAIS]


def _descricao_ofertada(produto: Produto, exigencias: list[Exigencia]) -> str:
    """Descrição para a proposta: só nome e atributos do catálogo que o motor de regras confirmou."""
    ident = f"marca {produto.marca or PREENCHER}, modelo {produto.modelo or PREENCHER}"
    fabricante = f", fabricante {produto.fabricante}" if produto.fabricante else ""
    tecnicos = _tecnicos(exigencias)
    return f"{produto.nome} — {ident}{fabricante}." + (f" {'; '.join(tecnicos)}." if tecnicos else "")


def _documentos(detalhes: list[tuple[DetalheRequisito, Any]], ofertavel: bool) -> list[DocumentoTecnico]:
    docs = []
    for d, req in detalhes:
        if natureza(d.atributo) != "documental":
            continue
        nome = rotulo(d.atributo)
        exigido = bool(req and req.obrigatorio)
        ev = req.evidencia_id if req else None
        if d.resultado == ResultadoRequisito.ATENDE:
            docs.append(DocumentoTecnico(nome=nome, situacao="declarado", detalhe=f"o catálogo declara: {d.valor_produto}", exigido=exigido, evidencia_id=ev))
        elif d.resultado == ResultadoRequisito.NAO_ATENDE:
            docs.append(DocumentoTecnico(nome=nome, situacao="falta", detalhe="o edital pede e o catálogo não declara: providenciar", exigido=exigido, evidencia_id=ev))
        elif d.atributo == "doc_amostra":
            docs.append(DocumentoTecnico(nome=nome, situacao="operacional", detalhe="o edital prevê amostra: separar a peça e acompanhar o prazo", exigido=exigido, evidencia_id=ev))
        else:
            docs.append(DocumentoTecnico(nome=nome, situacao="a_confirmar", detalhe=d.justificativa, exigido=exigido, evidencia_id=ev))
    if ofertavel:
        docs.append(DocumentoTecnico(nome="Ficha técnica do fabricante", situacao="recomendado",
                                     detalhe="comprova marca, modelo e os atributos da descrição", exigido=False))
    return docs


def _item(item: ItemDemanda, a: Analise, produtos: dict[str, Produto], reqs: dict[str, Any]) -> ItemProposta:
    base = {
        "item_id": item.id,
        "numero": item.numero_item,
        "descricao_orgao": item.descricao,
        "quantidade": item.quantidade,
        "unidade": item.unidade,
        "valor_unit_estimado": item.valor_unit_estimado,
        "valor_total_estimado": (item.quantidade or 0) * item.valor_unit_estimado if item.valor_unit_estimado is not None else None,
    }
    melhor = _melhor([c for c in a.correspondencias if c.item_id == item.id])
    if melhor is None:
        return ItemProposta(**base, situacao=Situacao.FORA_DO_CATALOGO, rotulo_situacao=ROTULO_SITUACAO[Situacao.FORA_DO_CATALOGO],
                            motivo="nenhum produto do catálogo nesta categoria")

    detalhes = [(d, reqs.get(d.requisito_id)) for d in melhor.detalhes]
    exigencias = [
        Exigencia(
            atributo=d.atributo, rotulo=rotulo(d.atributo), obrigatorio=bool(req.obrigatorio) if req else True,
            exigido=_exigido(req.operador.value if req else None, d.valor_exigido, d.atributo),
            ofertado=valor_tecnico(d.valor_produto, unidade(d.atributo)), resultado=d.resultado, justificativa=d.justificativa,
            evidencia_id=req.evidencia_id if req else None,
        )
        for d, req in detalhes
    ]
    p = produtos.get(melhor.produto_id)
    produto = ProdutoOfertado(sku=p.sku, nome=p.nome, fabricante=p.fabricante, marca=p.marca, modelo=p.modelo) if p else \
        ProdutoOfertado(sku=melhor.produto_id, nome="(produto fora do catálogo ativo)")

    def pend(e: Exigencia) -> str:
        texto = e.justificativa.removeprefix(f"{e.rotulo}: ")  # a justificativa do motor já começa pelo rótulo
        return f"{e.rotulo}{'' if e.obrigatorio else ' (desejável)'}: {texto}"

    if melhor.resultado == ResultadoItem.INCOMPATIVEL:
        motivos = [e.justificativa for e in exigencias if e.obrigatorio and e.resultado == ResultadoRequisito.NAO_ATENDE]
        return ItemProposta(**base, situacao=Situacao.NAO_OFERTAR, rotulo_situacao=ROTULO_SITUACAO[Situacao.NAO_OFERTAR],
                            motivo=f"o melhor candidato ({produto.sku}) não atende: " + "; ".join(motivos[:3]), exigencias=exigencias)

    if melhor.resultado == ResultadoItem.COMPATIVEL:
        situacao, motivo = Situacao.OFERTAR, "todas as exigências obrigatórias atendidas"
        pendencias = [pend(e) for e in exigencias if e.resultado != ResultadoRequisito.ATENDE]
    elif melhor.resultado == ResultadoItem.PARCIAL:
        situacao, motivo = Situacao.OFERTAR_COM_RESSALVA, "obrigatórias atendidas; há exigência desejável não atendida"
        pendencias = [pend(e) for e in exigencias if e.resultado != ResultadoRequisito.ATENDE]
    else:
        situacao = Situacao.RESOLVER_ANTES
        motivo = "falta informação para confirmar o atendimento"
        pendencias = [f"{pend(e)} Confirme antes de ofertar." for e in exigencias if e.resultado == ResultadoRequisito.INDETERMINADO and e.obrigatorio]
        pendencias += list(melhor.observacoes)
    if p and not (p.marca and p.modelo):
        pendencias.append("O catálogo não informa marca e modelo deste produto; os editais pedem os dois na proposta.")

    return ItemProposta(
        **base, situacao=situacao, rotulo_situacao=ROTULO_SITUACAO[situacao], motivo=motivo, produto=produto,
        descricao_ofertada=_descricao_ofertada(p, exigencias) if p else None, exigencias=exigencias, pendencias=pendencias,
        documentos=_documentos(detalhes, situacao in OFERTAVEIS),
    )


def _condicoes(a: Analise) -> list[Condicao]:
    c = a.contratacao
    q = a.qualificacao
    condicoes: list[Condicao] = []
    if c.data_abertura_proposta:
        condicoes.append(Condicao(rotulo="Início do recebimento de propostas", valor=f"{c.data_abertura_proposta:%d/%m/%Y %H:%M}", fonte="PNCP"))
    if c.data_encerramento_proposta:
        estado = f" ({q.prazo_status.lower()})" if q and q.prazo_status in {"ABERTA", "ENCERRADA"} else ""
        condicoes.append(Condicao(rotulo="Fim do recebimento de propostas", valor=f"{c.data_encerramento_proposta:%d/%m/%Y %H:%M}{estado}", fonte="PNCP"))
    # Exigências gerais do edital: a primeira ocorrência com prova, preferindo a achada pelo Investigador
    for atributo, nome in (("prazo_entrega_dias", "Prazo de entrega exigido"), ("garantia_anos", "Garantia mínima exigida"), ("doc_amostra", "Amostra")):
        cand = sorted((r for r in a.requisitos if r.atributo == atributo and r.evidencia_id), key=lambda r: r.fonte_extracao != "investigador")
        if not cand:
            continue
        r = cand[0]
        valor = "o edital prevê apresentação de amostra" if atributo == "doc_amostra" else _exigido(r.operador.value, r.valor, atributo)
        condicoes.append(Condicao(rotulo=nome, valor=valor, fonte="edital", evidencia_id=r.evidencia_id))
    return condicoes


def montar_kit(a: Analise, empresa: Empresa) -> KitProposta:
    produtos = {p.id: p for p in empresa.produtos}
    reqs = {r.id: r for r in a.requisitos}
    itens = [_item(i, a, produtos, reqs) for i in sorted(a.itens, key=lambda i: i.numero_item)]
    por_situacao = {s.value: sum(1 for i in itens if i.situacao == s) for s in Situacao}
    ofertaveis = [i for i in itens if i.situacao in OFERTAVEIS]
    c = a.contratacao
    edital = [{"id": d.id, "nome": d.nome, "tipo": d.tipo_doc} for d in a.documentos
              if d.caminho_local and not d.id.startswith("pncp-item") and "compactado" not in d.nome]

    avisos = [
        "Rascunho para revisão: o ObraNexo não envia propostas. Quem revisa e envia é o fornecedor.",
        "Sem preço: o valor estimado pelo órgão aparece só como referência; o preço é decisão do fornecedor.",
    ]
    if any(i.produto and not (i.produto.marca and i.produto.modelo) for i in ofertaveis):
        avisos.append("Há produtos sem marca ou modelo no catálogo: preencha antes de enviar, porque os editais pedem os dois.")

    return KitProposta(
        analise_id=a.id,
        gerado_em=datetime.now(),
        contratacao={
            "id": c.id, "orgao": c.orgao_nome, "municipio": c.municipio, "uf": c.uf, "modalidade": c.modalidade,
            "numero_compra": c.numero_compra, "ano": c.ano, "objeto": c.objeto, "valor_estimado": c.valor_estimado,
            "srp": c.srp, "link_oficial": c.link_oficial,
        },
        fornecedor={"nome": empresa.nome, "cnpj": empresa.cnpj},
        resumo={
            "itens": len(itens),
            "por_situacao": por_situacao,
            "ofertaveis": len(ofertaveis),
            "valor_estimado_ofertavel": round(sum(i.valor_total_estimado or 0 for i in ofertaveis), 2),
        },
        condicoes=_condicoes(a),
        itens=itens,
        habilitacao={
            "situacao": "em_construcao",
            "texto": "A lista de documentos de habilitação e de declarações será extraída do edital, com prova, na próxima etapa. "
                     "Até lá, confira as seções de habilitação no edital.",
            "documentos_do_edital": edital,
        },
        avisos=avisos,
    )


# ---------------------------------------------------------------------------------------------- exportação


def _fonte(a: Analise, evidencia_id: str | None) -> str:
    ev = a.evidencia(evidencia_id)
    if not ev:
        return ""
    if ev.documento_id.startswith("pncp-item"):
        return f"descrição do item no PNCP: “{ev.trecho_literal[:160]}”"
    doc = next((d for d in a.documentos if d.id == ev.documento_id), None)
    return f"{doc.nome if doc else ev.documento_id}, p. {ev.pagina}: “{ev.trecho_literal[:160]}”"


def kit_markdown(kit: KitProposta, a: Analise) -> str:
    c, r = kit.contratacao, kit.resumo
    linhas = [
        f"# Kit da proposta — {c['orgao']}",
        "",
        f"**Contratação:** {c['id']} · {c.get('modalidade') or ''} nº {c.get('numero_compra') or '—'}/{c.get('ano') or ''}  ",
        f"**Objeto:** {c['objeto']}  ",
        f"**Fornecedor:** {kit.fornecedor['nome']} · CNPJ {kit.fornecedor.get('cnpj') or PREENCHER}  ",
        f"**Fonte oficial:** {c.get('link_oficial') or '—'}",
        "",
        *[f"> {x}" for x in kit.avisos],
        "",
        f"**{r['ofertaveis']} de {r['itens']} itens ofertáveis** · valor estimado pelo órgão nesses itens (referência): "
        f"R$ {_num(r['valor_estimado_ofertavel'])}",
        "",
        "## Prazos e condições",
        "",
    ]
    for cd in kit.condicoes:
        fonte = _fonte(a, cd.evidencia_id)
        linhas.append(f"- **{cd.rotulo}:** {cd.valor}" + (f" — {fonte}" if fonte else f" ({cd.fonte})"))
    ofertaveis = [i for i in kit.itens if i.situacao in OFERTAVEIS or i.situacao == Situacao.RESOLVER_ANTES]
    linhas += ["", "## Itens da proposta", "", "| Item | Situação | Qtd | Produto | Marca | Modelo | Valor unit. estimado (ref.) |", "|---|---|---|---|---|---|---|"]
    for i in ofertaveis:
        p = i.produto
        linhas.append(f"| {i.numero} | {i.rotulo_situacao} | {valor_tecnico(i.quantidade)} {i.unidade or ''} | {p.sku if p else '—'} | "
                      f"{(p.marca if p else None) or PREENCHER} | {(p.modelo if p else None) or PREENCHER} | "
                      f"{'R$ ' + _num(i.valor_unit_estimado) if i.valor_unit_estimado is not None else '—'} |")
    for i in ofertaveis:
        linhas += ["", f"### Item {i.numero} — {i.rotulo_situacao}", "", f"**Pedido do órgão:** {i.descricao_orgao[:400]}", ""]
        if i.descricao_ofertada:
            linhas += [f"**Descrição para a proposta:** {i.descricao_ofertada}", ""]
        linhas += ["| Exigência | Edital pede | Ofertado | Resultado | Prova |", "|---|---|---|---|---|"]
        for e in i.exigencias:
            linhas.append(f"| {e.rotulo}{'' if e.obrigatorio else ' (desejável)'} | {e.exigido} | {e.ofertado} | {e.resultado.value} | "
                          f"{_fonte(a, e.evidencia_id).replace('|', '/') or '—'} |")
        if i.documentos:
            linhas += ["", "**Documentos técnicos:**"] + [f"- {d.nome}: {d.detalhe}" for d in i.documentos]
        if i.pendencias:
            linhas += ["", "**Resolver antes de enviar:**"] + [f"- {x}" for x in i.pendencias]
    fora = [i for i in kit.itens if i.situacao in {Situacao.NAO_OFERTAR, Situacao.FORA_DO_CATALOGO}]
    if fora:
        linhas += ["", "## Itens que não entram na proposta", ""] + [f"- **Item {i.numero}** ({i.rotulo_situacao}): {i.motivo}" for i in fora]
    linhas += ["", "## Habilitação e declarações", "", kit.habilitacao["texto"], ""]
    linhas += [f"- {d['nome']}" + (f" ({d['tipo']})" if d.get("tipo") else "") for d in kit.habilitacao["documentos_do_edital"]]
    linhas += ["", "---", f"*Gerado pelo ObraNexo v{kit.versao} em {kit.gerado_em:%d/%m/%Y %H:%M} a partir da análise {kit.analise_id}. "
               "Apoio à decisão: a proposta e o envio são de responsabilidade do fornecedor.*", ""]
    return "\n".join(linhas)


def kit_planilha(kit: KitProposta) -> bytes:
    """Planilha para preencher a proposta no portal de compras: uma linha por item, prazos numa segunda aba."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Proposta"
    colunas = [
        ("Item", 7), ("Situação", 22), ("Descrição do item (edital)", 50), ("Quantidade", 12), ("Unidade", 12),
        ("Valor unitário estimado (referência)", 18), ("Valor total estimado (referência)", 18), ("Fabricante", 30), ("Marca", 18),
        ("Modelo", 22), ("Descrição para a proposta", 70), ("Resolver antes de enviar", 50), ("Documentos técnicos", 50),
    ]
    ws.append([nome for nome, _ in colunas])
    for i in kit.itens:
        p = i.produto if i.situacao not in {Situacao.NAO_OFERTAR, Situacao.FORA_DO_CATALOGO} else None
        ws.append([
            i.numero, i.rotulo_situacao, i.descricao_orgao, i.quantidade, i.unidade, i.valor_unit_estimado, i.valor_total_estimado,
            (p.fabricante if p else None) or ("" if not p else PREENCHER), (p.marca if p else None) or ("" if not p else PREENCHER),
            (p.modelo if p else None) or ("" if not p else PREENCHER), i.descricao_ofertada or i.motivo,
            "\n".join(i.pendencias), "\n".join(f"{d.nome}: {d.detalhe}" for d in i.documentos),
        ])
    cabecalho = PatternFill("solid", start_color="0B2A3C")
    for cel in ws[1]:
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = cabecalho
        cel.alignment = Alignment(wrap_text=True, vertical="center")
    for n, (_, largura) in enumerate(colunas, start=1):
        ws.column_dimensions[ws.cell(row=1, column=n).column_letter].width = largura
    for linha in ws.iter_rows(min_row=2):
        for cel in linha:
            cel.alignment = Alignment(wrap_text=True, vertical="top")
        for cel in (linha[5], linha[6]):
            cel.number_format = '"R$" #,##0.00'
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    cond = wb.create_sheet("Prazos e condições")
    cond.append(["Condição", "Valor", "Fonte"])
    for cd in kit.condicoes:
        cond.append([cd.rotulo, cd.valor, cd.fonte])
    for cel in cond[1]:
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = cabecalho
    cond.column_dimensions["A"].width, cond.column_dimensions["B"].width, cond.column_dimensions["C"].width = 36, 40, 12

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------------------------- proposta de preços (.docx)


def _condicao(kit: KitProposta, rotulo_: str) -> str | None:
    return next((c.valor for c in kit.condicoes if c.rotulo == rotulo_), None)


def proposta_docx(kit: KitProposta, empresa: Empresa) -> bytes:
    """Proposta de preços pronta para revisar, precificar e assinar: modelo padrão com os itens ofertáveis.

    Campos que o ObraNexo não conhece (preços, validade, local e data, dados ausentes do fornecedor) saem
    como "[preencher]" destacado em amarelo. A última página é uma conferência interna (pendências, itens fora
    da proposta) e deve ser removida antes do envio."""
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.8)
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(10.5)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Arial")
    normal.paragraph_format.space_after = Pt(4)

    def campo(par, valor: str | None, negrito: bool = False, tamanho: float | None = None):
        run = par.add_run(valor if valor else PREENCHER)
        run.bold = negrito
        if tamanho:
            run.font.size = Pt(tamanho)
        if not valor:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        return run

    def linha(rotulo_: str, valor: str | None):
        p = doc.add_paragraph()
        p.add_run(f"{rotulo_}: ").bold = True
        campo(p, valor)
        return p

    def titulo(texto: str, tamanho: float = 11.5, centro: bool = False):
        p = doc.add_paragraph()
        r = p.add_run(texto)
        r.bold, r.font.size = True, Pt(tamanho)
        p.paragraph_format.space_before = Pt(12)
        if centro:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return p

    def sombrear(celula, cor: str):
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), cor)
        celula._tc.get_or_add_tcPr().append(shd)

    c = kit.contratacao
    # Cabeçalho do fornecedor
    p = doc.add_paragraph()
    campo(p, empresa.razao_social or empresa.nome, negrito=True, tamanho=13)
    linha("CNPJ", empresa.cnpj)
    linha("Endereço", empresa.endereco)
    p = doc.add_paragraph()
    p.add_run("Telefone: ").bold = True
    campo(p, empresa.telefone)
    p.add_run(" · E-mail: ").bold = True
    campo(p, empresa.email)

    # Destinatário e referência
    doc.add_paragraph()
    doc.add_paragraph().add_run(f"Ao(À) {c['orgao']}").bold = True
    doc.add_paragraph("A/C: Agente de contratação / Pregoeiro(a) e equipe de apoio")
    referencia = f"{c.get('modalidade') or 'Licitação'} nº {c.get('numero_compra') or '—'}/{c.get('ano') or ''}"
    linha("Referência", f"{referencia} · controle PNCP {c['id']}")
    linha("Objeto", c.get("objeto"))

    titulo("PROPOSTA DE PREÇOS", 14, centro=True)
    doc.add_paragraph("Apresentamos nossa proposta para o fornecimento dos itens abaixo, em conformidade com as especificações "
                      "e as condições do edital e de seus anexos.")

    # Itens ofertados
    itens = [i for i in kit.itens if i.situacao in OFERTAVEIS and i.produto]
    cab = ["Item", "Especificação do produto ofertado", "Marca / modelo / fabricante", "Unid.", "Qtd.", "Valor unitário (R$)", "Valor total (R$)"]
    larguras = [1.1, 7.0, 3.6, 1.4, 1.5, 1.9, 1.9]
    tabela = doc.add_table(rows=1, cols=len(cab))
    tabela.style = "Table Grid"
    tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
    for n, texto in enumerate(cab):
        cel = tabela.rows[0].cells[n]
        r = cel.paragraphs[0].add_run(texto)
        r.bold, r.font.size, r.font.color.rgb = True, Pt(9), RGBColor(0xFF, 0xFF, 0xFF)
        sombrear(cel, "0B2A3C")
    for i in itens:
        celulas = tabela.add_row().cells
        prod = i.produto
        tecnicos = _tecnicos(i.exigencias)
        celulas[0].text = str(i.numero)
        celulas[1].text = f"{prod.nome}." + (f" {'; '.join(tecnicos)}." if tecnicos else "")
        par = celulas[2].paragraphs[0]
        for k, (rot, val) in enumerate((("Marca", prod.marca), ("Modelo", prod.modelo), ("Fabricante", prod.fabricante))):
            if k:
                par.add_run().add_break()
            par.add_run(f"{rot}: ")
            campo(par, val)
        celulas[3].text = i.unidade or ""
        celulas[4].text = valor_tecnico(i.quantidade) if i.quantidade is not None else ""
        campo(celulas[5].paragraphs[0], None)
        campo(celulas[6].paragraphs[0], None)
    total = tabela.add_row().cells
    rotulo_total = total[0].merge(total[5])
    rotulo_total.paragraphs[0].add_run("VALOR GLOBAL DA PROPOSTA (R$)").bold = True
    rotulo_total.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    campo(total[6].paragraphs[0], None, negrito=True)
    for linha_ in tabela.rows:
        for n, cel in enumerate(linha_.cells):
            cel.width = Cm(larguras[n])
            for par in cel.paragraphs:
                for run in par.runs:
                    run.font.size = run.font.size or Pt(9)

    # Condições
    titulo("CONDIÇÕES DA PROPOSTA")
    prazo = _condicao(kit, "Prazo de entrega exigido")
    garantia = _condicao(kit, "Garantia mínima exigida")
    linha("Validade da proposta", None)
    linha("Prazo de entrega", f"{prazo.replace('máximo', 'até')}, conforme o edital" if prazo else None)
    linha("Garantia", f"conforme o edital ({garantia.replace('mínimo', 'mínimo de')})" if garantia else None)
    linha("Local de entrega", "conforme o edital")
    linha("Dados bancários", empresa.dados_bancarios)

    titulo("DECLARAÇÕES")
    declaracoes = (
        "Declaramos que os produtos ofertados atendem às especificações técnicas exigidas no edital e em seus anexos.",
        "Declaramos que nos preços propostos estão incluídas todas as despesas com tributos, fretes, seguros e demais encargos "
        "necessários ao fornecimento.",
        "Declaramos conhecer e aceitar todas as condições do edital, inclusive os prazos de entrega e de garantia.",
    )
    for n, texto in enumerate(declaracoes, start=1):
        doc.add_paragraph(f"{n}. {texto}")

    # Assinatura: o bloco inteiro fica na mesma página ("manter com o próximo" em cada parágrafo)
    bloco = [doc.add_paragraph(), linha("Local e data", None), doc.add_paragraph(), doc.add_paragraph()]
    traco = doc.add_paragraph("_______________________________________________")
    traco.alignment = WD_ALIGN_PARAGRAPH.CENTER
    bloco.append(traco)
    for valor, negrito in ((empresa.representante, True), (empresa.representante_cargo, False), (empresa.razao_social or empresa.nome, False)):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        campo(p, valor, negrito=negrito)
        bloco.append(p)
    for p in bloco:
        p.paragraph_format.keep_with_next = True
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("CNPJ ")
    campo(p, empresa.cnpj)

    # Conferência interna: nunca vai para o órgão
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    p = titulo("CONFERÊNCIA INTERNA — REMOVA ESTA PÁGINA ANTES DE ENVIAR", 12)
    p.runs[0].font.highlight_color = WD_COLOR_INDEX.YELLOW
    doc.add_paragraph(f"Rascunho gerado pelo ObraNexo em {kit.gerado_em:%d/%m/%Y %H:%M} a partir da análise {kit.analise_id}. "
                      "Confira com o edital, preencha os campos destacados em amarelo e defina os preços. "
                      "Se o edital trouxer um modelo próprio de proposta, transfira estes dados para ele.")
    com_pendencia = [i for i in itens if i.pendencias]
    if com_pendencia:
        titulo("Resolver antes de assinar", 11)
        for i in com_pendencia:
            for x in i.pendencias:
                doc.add_paragraph(f"Item {i.numero}: {x}", style="List Bullet")
    fora = [i for i in kit.itens if i not in itens]
    if fora:
        titulo("Itens que não entraram na proposta", 11)
        for i in fora:
            detalhe = "; ".join(i.pendencias) if i.situacao == Situacao.RESOLVER_ANTES and i.pendencias else i.motivo
            doc.add_paragraph(f"Item {i.numero} ({i.rotulo_situacao}): {detalhe}", style="List Bullet")
    if kit.habilitacao["documentos_do_edital"]:
        titulo("Documentos da licitação para conferir (habilitação, declarações, modelo de proposta)", 11)
        for d in kit.habilitacao["documentos_do_edital"]:
            doc.add_paragraph(d["nome"], style="List Bullet")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
